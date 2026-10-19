"""Stage 2 outputs: exports/analysis/<SYMBOL>.md (readable report), <SYMBOL>.xlsx (the
Module 13 model laid out sheet by sheet) and exports/analysis_summary.csv."""

from __future__ import annotations

import csv
import math

from ..config import EXPORT_DIR

SECTIONS = [("P&L", ["revenue", "growth", "ebitda", "da", "ebit", "interest", "pbt", "tax", "pat", "dividends", "eps"]),
            ("Balance sheet", ["net_block", "inventory", "receivables", "cash", "other_assets", "total_assets",
                               "payables", "other_liabilities", "debt", "equity", "total_liabilities_equity"]),
            ("Cash flow", ["cfo", "capex", "cfi", "cff", "delta_nwc", "fcff"])]
LABEL = {"revenue": "Revenue", "growth": "Revenue growth", "ebitda": "EBITDA", "da": "Depreciation & amortisation",
         "ebit": "EBIT", "interest": "Interest", "pbt": "Profit before tax", "tax": "Tax", "pat": "Profit after tax",
         "dividends": "Dividends", "eps": "EPS", "net_block": "Net block (PP&E)", "inventory": "Inventory",
         "receivables": "Receivables", "cash": "Cash", "other_assets": "Other assets (held)",
         "total_assets": "Total assets", "payables": "Payables", "other_liabilities": "Other liabilities (held)",
         "debt": "Debt", "equity": "Shareholders' equity", "total_liabilities_equity": "Total liabilities & equity",
         "cfo": "Cash from operations", "capex": "Capex", "cfi": "Cash from investing", "cff": "Cash from financing",
         "delta_nwc": "Change in net working capital", "fcff": "Free cash flow to the firm"}
PCT_LINES = {"growth"}


def _p(v, d=1):
    return "—" if v is None or (isinstance(v, float) and not math.isfinite(v)) else f"{v * 100:+.{d}f}%"


def _n(v, scale=1e6, d=0):
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return "—"
    return f"{v / scale:,.{d}f}"


def _money(v, d=2):
    return "—" if v is None else f"{v:,.{d}f}"


def markdown(o: dict) -> str:
    if o.get("error"):
        return f"# {o.get('name', o['symbol'])}\n\nNot analysed: {o['error']}\n"
    c, cc, v = o["classification"], o["cost_of_capital"], o["valuations"]
    ccy, sc = o["currency"], o["statement_ccy"]
    out = [f"# {o['name']} ({o['symbol']})\n",
           f"As of {o['as_of']}. Price {_money(o['price'])} {ccy}. Statements in {sc}"
           + (f" (1 {sc} = {o['fx_trade_per_stmt']:,.4f} {ccy})" if sc != ccy else "") + ".\n",
           f"**Verdict: {o['verdict']}** — combined fair value {_money(o['fair_value'])} {ccy} "
           f"({_p(o['upside'])} vs price).\n"]
    out.append("| Method | Fair value per share | vs price |")
    out.append("|---|---|---|")
    names = {"dcf": "DCF (FCFF, Module 13)", "bank_pb": "Justified P/B (financials)", "relative": "Peer multiples (median)"}
    for m, lbl in names.items():
        if m in v and v[m].get("per_share_trade"):
            out.append(f"| {lbl} | {_money(v[m]['per_share_trade'])} {ccy} | {_p(v[m]['per_share_trade'] / o['price'] - 1)} |")
    out.append("")
    if o["flags"]:
        out.append("Flags:\n" + "\n".join(f"- {f}" for f in o["flags"]) + "\n")
    out.append("## Stage 0 categories\n")
    out.append("| Market (MSCI / FTSE) | Listing | Sector group | Cap band | Style | Behaviour | News (Stage 1) |")
    out.append("|---|---|---|---|---|---|---|")
    out.append(f"| {c['market_class']} / {c['market_class_ftse']} | {c['listing_structure']} | {c['sector_group']} | "
               f"{c['cap_band']}{' (SEBI ' + c['cap_band_sebi'] + ')' if c.get('cap_band_sebi') else ''} | {c['growth_value']} | "
               f"{c['cyclical_defensive']} | {c['news_verdict'] or '—'} ({c['news_composite'] or '—'}) |\n")
    out.append("## Growth\n")
    out.append("| Horizon | Expected | 5–95% range | Chance of beating cash | Source |")
    out.append("|---|---|---|---|---|")
    for h, g in (o.get("growth") or {}).items():
        lo, hi = (g.get("range_5_95") or [None, None])
        pbc = g.get("p_beat_cash")
        pbc_s = "—" if pbc is None else f"{pbc:.0%}"
        out.append(f"| {h} | {_p(g['expected'])} | {_p(lo, 0)} to {_p(hi, 0)} | {pbc_s} | {g.get('model')} (Milestone 2) |")
    lt = o["long_term_growth"]
    out.append(f"| 5 years (Module 13 projection) | revenue {_p(lt['revenue_cagr_5y'])} a year, EPS {_p(lt['eps_cagr_5y'])} a year | | | integrated model |\n")
    out.append("## Cost of capital (Module 13 Ch 16–17)\n")
    out.append("| Item | Value | Source |")
    out.append("|---|---|---|")
    for k, lbl, src in (("rf", "Risk-free (10-year yield)", "rf_source"), ("erp", "Equity risk premium", "erp_source"),
                        ("beta", "Beta", "beta_source"), ("ke", "Cost of equity = Rf + β·ERP", None),
                        ("kd_pre_tax", "Cost of debt (pre-tax)", None), ("kd_after_tax", "Cost of debt after tax shield", None),
                        ("tax_rate", "Tax rate", None), ("debt_weight", "Debt weight (market values)", None),
                        ("wacc", "WACC", None), ("g", "Terminal growth", "g_source")):
        val = cc[k] if k == "beta" else _p(cc[k], 2).lstrip("+")
        out.append(f"| {lbl} | {val if k != 'beta' else f'{val:.2f}'} | {cc.get(src, '') if src else ''} |")
    out.append("")
    out.append("## Assumptions (Module 13 Ch 4–5)\n")
    out.append("| Assumption | Value | Basis |")
    out.append("|---|---|---|")
    for name, val, basis in o["assumptions"]:
        shown = f"{val:.1f}" if "days" in name.lower() else (f"{val:.2f}×" if "turnover" in name.lower() else _p(val).lstrip("+"))
        out.append(f"| {name} | {shown} | {basis} |")
    out.append("")
    pr = o["projection"]
    yrs = pr["years"]
    out.append(f"## Integrated model ({sc} millions; years ending {', '.join(y[:4] for y in yrs[:pr['n_hist']])} reported, "
               f"then projected)\n")
    for title, lines in SECTIONS:
        out.append(f"**{title}**\n")
        out.append("| | " + " | ".join(y[:4] if not y.endswith("E") else y for y in yrs) + " |")
        out.append("|---|" + "---|" * len(yrs))
        for ln in lines:
            vals = pr["rows"][ln]
            cells = [_p(x) if ln in PCT_LINES else (_money(x) if ln == "eps" else _n(x)) for x in vals]
            out.append(f"| {LABEL[ln]} | " + " | ".join(cells) + " |")
        out.append("")
    bal = max(abs(x) for x in pr["balance_diff"]) if pr["balance_diff"] else 0
    out.append(f"Balance check: largest projected difference between assets and liabilities + equity = {bal:,.2f} {sc}.\n")
    if "dcf" in v:
        d = v["dcf"]
        out.append("## DCF (Module 13 Ch 15–18)\n")
        n = len(d["pv_fcff"])
        fc = pr["rows"]["fcff"][-n:]
        out.append("| | " + " | ".join(yrs[-n:]) + " |")
        out.append("|---|" + "---|" * n)
        out.append("| FCFF | " + " | ".join(_n(x) for x in fc) + " |")
        out.append("| Discount factor | " + " | ".join(f"{x:.3f}" for x in d["discount_factors"]) + " |")
        out.append("| Present value | " + " | ".join(_n(x) for x in d["pv_fcff"]) + " |")
        out.append("")
        out.append(f"Terminal value {_n(d['terminal_value'])}m (present value {_n(d['pv_terminal'])}m, "
                   f"{d['terminal_share_of_ev']:.0%} of EV). Enterprise value {_n(d['ev'])}m − debt + cash − minority = "
                   f"equity {_n(d['equity_value'])}m ÷ {o['shares'] / 1e6:,.1f}m shares = {_money(d['per_share'])} {sc} "
                   f"= **{_money(d['per_share_trade'])} {ccy}** per share.\n")
        grid = d["sensitivity"]
        gs = list(next(iter(grid.values())).keys())
        out.append(f"Sensitivity: fair value per share ({ccy}), WACC down the side, terminal growth across\n")
        out.append("| WACC \\ g | " + " | ".join(f"{g * 100:.1f}%" for g in gs) + " |")
        out.append("|---|" + "---|" * len(gs))
        for wk, row in grid.items():
            out.append(f"| {wk * 100:.1f}% | " + " | ".join(_money(x, 0) for x in row.values()) + " |")
        out.append("")
    if "bank_pb" in v:
        b = v["bank_pb"]
        out.append("## Justified price-to-book (financials)\n")
        out.append(f"ROE {_p(b['roe']).lstrip('+')} (3-year average), Ke {_p(cc['ke'], 2).lstrip('+')}, g {_p(cc['g'], 2).lstrip('+')} → "
                   f"justified P/B {b['justified_pb']:.2f}× book value {_n(b['book_value'])}m = "
                   f"**{_money(b['per_share_trade'])} {ccy}** per share.\n")
    if o.get("peers"):
        out.append("## Peer multiples (Module 13 Ch 14, relative valuation)\n")
        out.append("| Company | P/E | EV/EBITDA | P/B | EV/Sales |")
        out.append("|---|---|---|---|---|")
        fm = lambda x: "—" if x is None else f"{x:.1f}"
        own = o.get("own_multiples") or {}
        out.append(f"| **{o['name']}** | {fm(own.get('pe'))} | {fm(own.get('ev_ebitda'))} | {fm(own.get('pb'))} | {fm(own.get('ev_sales'))} |")
        for p in o["peers"]:
            out.append(f"| {p['name']} | {fm(p.get('pe'))} | {fm(p.get('ev_ebitda'))} | {fm(p.get('pb'))} | {fm(p.get('ev_sales'))} |")
        if "relative" in v:
            r = v["relative"]
            implied = ", ".join(f"{k} {_money(x['per_share'] * o['fx_trade_per_stmt'])}" for k, x in r.items()
                                if isinstance(x, dict))
            out.append(f"\nImplied per share ({ccy}): {implied}; median **{_money(r['per_share_trade'])}**.\n")
    if "option_merton" in v:
        m = v["option_merton"]
        out.append("## Option-based check (Merton)\n")
        out.append(f"Implied asset value {_n(m['asset_value'])}m, asset volatility {m['asset_vol']:.0%}, distance to default "
                   f"{m['distance_to_default']:.2f}, one-year default probability {m['default_prob_1y']:.1%}.\n")
    out.append("Built with the Module 13 method on Yahoo Finance statements. Research output, not investment advice.\n")
    return "\n".join(out)


def workbook(o: dict, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="DDE4EE")
    ws = wb.active
    ws.title = "Summary"
    rows = [("Company", o["name"]), ("Symbol", o["symbol"]), ("As of", o["as_of"]), ("Price", o["price"]),
            ("Currency", o["currency"]), ("Statement currency", o["statement_ccy"]),
            ("Trading units per statement unit", o["fx_trade_per_stmt"]), ("Fair value (combined)", o["fair_value"]),
            ("Upside", o["upside"]), ("Verdict", o["verdict"])]
    for m, x in o["valuations"].items():
        if x.get("per_share_trade"):
            rows.append((f"Fair value: {m}", x["per_share_trade"]))
    rows += [("Flag", f) for f in o["flags"]]
    rows += [(f"Category: {k}", val) for k, val in o["classification"].items()]
    for r in rows:
        ws.append(list(r))
    for c in ws["A"]:
        c.font = bold
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    a = wb.create_sheet("Assumptions")
    a.append(["Assumption", "Value", "Basis"])
    for r in o["assumptions"]:
        a.append(list(r))
    a.append([])
    a.append(["Cost of capital", "Value", "Source"])
    cc = o["cost_of_capital"]
    for k in ("rf", "erp", "beta", "ke", "kd_pre_tax", "kd_after_tax", "tax_rate", "debt_weight", "wacc", "g"):
        a.append([k, cc[k], cc.get(f"{k}_source", "")])
    a.column_dimensions["A"].width = 40
    a.column_dimensions["C"].width = 70
    pr = o["projection"]
    for title, lines in SECTIONS:
        s = wb.create_sheet(title.replace("&", "and"))
        s.append([f"{o['statement_ccy']}"] + pr["years"])
        for ln in lines:
            s.append([LABEL[ln]] + [x for x in pr["rows"][ln]])
        s.column_dimensions["A"].width = 34
        for cell in s[1]:
            cell.font, cell.fill = bold, head
    if "dcf" in o["valuations"]:
        d = o["valuations"]["dcf"]
        s = wb.create_sheet("DCF")
        n = len(d["pv_fcff"])
        s.append(["", *pr["years"][-n:]])
        s.append(["FCFF", *pr["rows"]["fcff"][-n:]])
        s.append(["Discount factor", *d["discount_factors"]])
        s.append(["Present value", *d["pv_fcff"]])
        s.append([])
        for k in ("terminal_value", "pv_terminal", "ev", "equity_value", "per_share", "per_share_trade",
                  "terminal_share_of_ev"):
            s.append([k, d[k]])
        s.append([])
        s.append(["Sensitivity (per share, trading currency): WACC down, g across"])
        grid = d["sensitivity"]
        gs = list(next(iter(grid.values())).keys())
        s.append(["WACC \\ g", *gs])
        for wk, row in grid.items():
            s.append([wk, *row.values()])
        s.column_dimensions["A"].width = 30
    if o.get("peers"):
        s = wb.create_sheet("Peers")
        s.append(["Company", "Symbol", "P/E", "EV/EBITDA", "P/B", "EV/Sales"])
        for p in o["peers"]:
            s.append([p["name"], p["symbol"], p.get("pe"), p.get("ev_ebitda"), p.get("pb"), p.get("ev_sales")])
        s.column_dimensions["A"].width = 40
    wb.save(path)


def summary_csv(results: list[dict], path) -> None:
    rows = []
    for o in results:
        if o.get("error"):
            rows.append({"symbol": o["symbol"], "name": o.get("name"), "error": o["error"]})
            continue
        v = o["valuations"]
        rows.append({"symbol": o["symbol"], "name": o["name"], "currency": o["currency"], "price": o["price"],
                     "fair_value": o["fair_value"], "upside": o["upside"], "verdict": o["verdict"],
                     "dcf": (v.get("dcf") or {}).get("per_share_trade"),
                     "bank_pb": (v.get("bank_pb") or {}).get("per_share_trade"),
                     "relative": (v.get("relative") or {}).get("per_share_trade"),
                     "wacc": o["cost_of_capital"]["wacc"], "ke": o["cost_of_capital"]["ke"],
                     "g": o["cost_of_capital"]["g"], "growth_1y": ((o.get("growth") or {}).get("1 year") or {}).get("expected"),
                     "revenue_cagr_5y": o["long_term_growth"]["revenue_cagr_5y"],
                     "market_class": o["classification"]["market_class"],
                     "sector_group": o["classification"]["sector_group"], "flags": " | ".join(o["flags"])})
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def write(o: dict, out_dir=None) -> dict:
    out_dir = out_dir or EXPORT_DIR / "analysis"
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = o["symbol"].replace("/", "_")
    (out_dir / f"{safe}.md").write_text(markdown(o), encoding="utf-8")
    if not o.get("error"):
        workbook(o, out_dir / f"{safe}.xlsx")
    return {"md": str(out_dir / f"{safe}.md")}

